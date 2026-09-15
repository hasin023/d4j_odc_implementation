"""Cross-model comparison report: Gemini vs a second model (e.g. gpt-5-mini) on
Alvee's 13-bug manual-analysis shortlist, plus a small reasoning-effort experiment.

Tab 1 "13-Bug Comparison": one row per bug, all 4 arms (scientific/few x
prefix/postfix) for both models side by side against the manual ground truth
already recorded in analysis_v2/manual_analysis/shortlist.csv — a model cell is
highlighted when it disagrees with ground truth.

Tab 2 "Accuracy by Arm": per-arm accuracy % for each model against ground truth,
computed from the same data as Tab 1 (not hardcoded), plus a grouped bar chart.

Tab 3 "Reasoning Effort": two small tables transcribed from an ad-hoc terminal
experiment this session (never written to a tracked file, so the raw numbers are
hand-entered here) - per-effort-level timing across whichever bugs were tested at
each level, and the one apples-to-apples 4-bug default-vs-high accuracy comparison,
with a callout for the counter-intuitive Chart_9 flip (default correct, high wrong).

Output lands under .dist/study/reports/<artifacts_root_name>/model_comparison/, and
each run appends one line to that root's reports.jsonl ledger.

Usage: python scripts/reports/generate_model_comparison_report.py
       --provider openai-compatible --model gpt-5-mini
       [--artifacts-root .dist/study/artifacts_v2] [--output <path>]
"""
from __future__ import annotations

import argparse
import csv
import json
import subprocess
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT))
from d4j_odc_pipeline.odc import model_slug  # noqa: E402

SHORTLIST_CSV = REPO_ROOT / "analysis_v2/manual_analysis/shortlist.csv"

HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_FONT = Font(bold=True)
HIGHLIGHT = PatternFill(start_color="FFF3B0", end_color="FFF3B0", fill_type="solid")
TITLE_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

ARMS = ["sci_pre", "sci_post", "few_pre", "few_post"]
ARM_LABELS = {"sci_pre": "Scientific Prefix", "sci_post": "Scientific Postfix", "few_pre": "Few Prefix", "few_post": "Few Postfix"}

# Reasoning-effort experiment data (2026-09-15 session) — never written to a
# tracked file (all /tmp scratch during exploration), transcribed here from the
# terminal results. See docs/study_execution_log.md for the session narrative.
EFFORT_TIMING = [
    # (bug, condition, effort, seconds)
    ("Chart_1", "few_pre", "default", 29.4),
    ("Chart_1", "few_pre", "minimal", 12.5),
    ("Chart_1", "sci_pre", "default", 50.1),
    ("Chart_1", "sci_post", "default", 19.5),
    ("Chart_1", "sci_pre", "minimal", 20.7),
    ("Chart_1", "sci_post", "minimal", 26.5),
    ("Chart_2", "sci_pre", "default", 118.3),
    ("Chart_2", "sci_post", "default", 80.2),
    ("Chart_2", "sci_pre", "minimal", 42.9),
    ("Chart_2", "sci_post", "minimal", 25.6),
    ("Chart_3", "sci_pre", "default", 94.0),
    ("Chart_3", "sci_post", "default", 59.3),
    ("Chart_3", "sci_pre", "minimal", 28.6),
    ("Chart_3", "sci_post", "minimal", 21.7),
    ("Chart_9", "few_pre", "high", 66.9),
    ("Lang_40", "few_pre", "high", 75.9),
    ("Math_104", "few_pre", "high", 82.0),
    ("Chart_17", "few_pre", "high", 67.9),
]

# The one apples-to-apples comparison: same 4 bugs, same arm (few_pre), default
# vs high effort. Default values come from the actual classification.json files
# (loaded below); high-effort values are hand-transcribed from the terminal
# results since that ad-hoc test never wrote into the tracked artifacts tree.
EFFORT_HIGH_RESULTS = {
    "Chart_9": "Algorithm/Method",
    "Lang_40": "Algorithm/Method",
    "Math_104": "Assignment/Initialization",
    "Chart_17": "Checking",
}


def git_sha() -> str | None:
    try:
        out = subprocess.run(["git", "rev-parse", "--short", "HEAD"], capture_output=True, text=True, timeout=5, cwd=REPO_ROOT)
    except (OSError, subprocess.SubprocessError):
        return None
    return out.stdout.strip() or None if out.returncode == 0 else None


def load_json(path: Path) -> dict | None:
    if not path.exists():
        return None
    return json.loads(path.read_text())


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def pct(n, total):
    return f"{(100.0 * n / total) if total else 0.0:.1f}%"


# ── Data loading ─────────────────────────────────────────────────────────────

def load_shortlist() -> list[dict]:
    rows = []
    with SHORTLIST_CSV.open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            rows.append(row)
    return sorted(rows, key=lambda r: int(r["rank"]))


def load_second_model_arms(root: Path, project: str, bug_id: int, provider: str, model: str) -> dict:
    tag_slug = model_slug(provider, model)
    sci = load_json(root / "prefix" / f"{project}_{bug_id}_prefix" / f"classification.scientific-open.{tag_slug}.json")
    sci_post = load_json(root / "postfix" / f"{project}_{bug_id}_postfix" / f"classification.scientific-open.{tag_slug}.json")
    few = load_json(root / "prefix" / f"{project}_{bug_id}_prefix" / f"classification.few-open.{tag_slug}.json")
    few_post = load_json(root / "postfix" / f"{project}_{bug_id}_postfix" / f"classification.few-open.{tag_slug}.json")
    missing = [n for n, v in [("sci_pre", sci), ("sci_post", sci_post), ("few_pre", few), ("few_post", few_post)] if v is None]
    if missing:
        raise FileNotFoundError(f"{project}_{bug_id}: missing second-model classification(s): {', '.join(missing)}")
    return {
        "sci_pre": sci.get("odc_type"),
        "sci_post": sci_post.get("odc_type"),
        "few_pre": few.get("odc_type"),
        "few_post": few_post.get("odc_type"),
    }


def load_records(root: Path, provider: str, model: str) -> list[dict]:
    records = []
    for row in load_shortlist():
        project, bug_id = row["project"], int(row["bug_number"])
        second = load_second_model_arms(root, project, bug_id, provider, model)
        records.append({
            "bug_id": row["bug_id"],
            "truth": row["manual_ground_truth"],
            "gemini": {
                "sci_pre": row["sci_prefix"],
                "sci_post": row["sci_postfix"],
                "few_pre": row["few_prefix"],
                "few_post": row["few_postfix"],
            },
            "second": second,
        })
    return records


# ── Tab 1: 13-Bug Comparison ────────────────────────────────────────────────

def build_comparison_tab(wb, records, model_label):
    ws = wb.create_sheet("13-Bug Comparison")
    headers = ["Bug", "Ground Truth"] + [f"Gemini {ARM_LABELS[a]}" for a in ARMS] + [f"{model_label} {ARM_LABELS[a]}" for a in ARMS]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    for r in records:
        row_vals = [r["bug_id"], r["truth"]]
        row_vals += [r["gemini"][a] for a in ARMS]
        row_vals += [r["second"][a] for a in ARMS]
        ws.append(row_vals)
        row_idx = ws.max_row
        for i, a in enumerate(ARMS):
            gemini_col = 3 + i
            if r["gemini"][a] != r["truth"]:
                ws.cell(row=row_idx, column=gemini_col).fill = HIGHLIGHT
            second_col = 3 + len(ARMS) + i
            if r["second"][a] != r["truth"]:
                ws.cell(row=row_idx, column=second_col).fill = HIGHLIGHT

    widths = [12, 24] + [20] * (2 * len(ARMS))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"


# ── Tab 2: Accuracy by Arm ──────────────────────────────────────────────────

def build_accuracy_tab(wb, records, model_label):
    ws = wb.create_sheet("Accuracy by Arm")
    ws.append(["Model"] + [ARM_LABELS[a] for a in ARMS])
    style_header(ws, 1, len(ARMS) + 1)

    for model_key, name in [("gemini", "Gemini"), ("second", model_label)]:
        row_vals = [name]
        for a in ARMS:
            correct = sum(1 for r in records if r[model_key][a] == r["truth"])
            row_vals.append(correct / len(records))
        ws.append(row_vals)
        row_idx = ws.max_row
        for c in range(2, len(ARMS) + 2):
            ws.cell(row=row_idx, column=c).number_format = "0%"

    ws.column_dimensions["A"].width = 16
    for i in range(2, len(ARMS) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 18

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = "Accuracy vs Manual Ground Truth (13 bugs)"
    chart.y_axis.title = "Accuracy"
    chart.y_axis.numFmt = "0%"
    data = Reference(ws, min_col=1, max_col=len(ARMS) + 1, min_row=1, max_row=3)
    cats = Reference(ws, min_col=2, max_col=len(ARMS) + 1, min_row=1, max_row=1)
    chart.add_data(data, titles_from_data=True, from_rows=True)
    chart.set_categories(cats)
    chart.height = 9
    chart.width = 18
    ws.add_chart(chart, "A6")


# ── Tab 3: Reasoning Effort ──────────────────────────────────────────────────

def build_reasoning_effort_tab(wb, records, model_label):
    ws = wb.create_sheet("Reasoning Effort")
    ws.append([f"Timing by reasoning effort level ({model_label}, ad-hoc pilot — not the same bugs at every level, see note)"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append(["Bug", "Condition", "Effort", "Seconds"])
    style_header(ws, 2, 4)
    for bug, cond, effort, secs in EFFORT_TIMING:
        ws.append([bug, cond, effort, secs])
    timing_end = ws.max_row

    ws.append([])
    ws.append(["Avg seconds/call by effort level (pooled across whatever bugs were tested at that level):"])
    avg_row_start = ws.max_row + 1
    by_effort: dict[str, list[float]] = {}
    for _, _, effort, secs in EFFORT_TIMING:
        by_effort.setdefault(effort, []).append(secs)
    ws.append(["Effort", "Avg Seconds", "N"])
    style_header(ws, ws.max_row, 3)
    for effort in ["minimal", "default", "high"]:
        vals = by_effort.get(effort, [])
        if vals:
            ws.append([effort, sum(vals) / len(vals), len(vals)])
    avg_row_end = ws.max_row

    chart = BarChart()
    chart.type = "col"
    chart.title = "Avg seconds/call by reasoning effort"
    chart.y_axis.title = "Seconds"
    data = Reference(ws, min_col=2, max_col=2, min_row=avg_row_start, max_row=avg_row_end)
    cats = Reference(ws, min_col=1, max_col=1, min_row=avg_row_start + 1, max_row=avg_row_end)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 8
    chart.width = 14
    ws.add_chart(chart, f"F{avg_row_start}")

    ws.append([])
    ws.append(["Default vs high effort, same 4 bugs, same arm (few, prefix):"])
    ws.cell(row=ws.max_row, column=1).font = TITLE_FONT
    ws.append(["Bug", "Ground Truth", "Default Effort", "High Effort"])
    style_header(ws, ws.max_row, 4)
    default_correct, high_correct = 0, 0
    by_id = {r["bug_id"]: r for r in records}
    for bug in ["Chart_9", "Lang_40", "Math_104", "Chart_17"]:
        r = by_id.get(bug)
        truth = r["truth"] if r else "?"
        default_val = r["second"]["few_pre"] if r else "?"
        high_val = EFFORT_HIGH_RESULTS[bug]
        ws.append([bug, truth, default_val, high_val])
        row_idx = ws.max_row
        if default_val == truth:
            default_correct += 1
        else:
            ws.cell(row=row_idx, column=3).fill = HIGHLIGHT
        if high_val == truth:
            high_correct += 1
        else:
            ws.cell(row=row_idx, column=4).fill = HIGHLIGHT

    ws.append([])
    callout = ws.cell(
        row=ws.max_row + 1, column=1,
        value=(
            f"Default effort scored {default_correct}/4 correct vs high effort's {high_correct}/4 on this "
            "4-bug apples-to-apples comparison. Chart_9 in particular flipped from correct (default) to "
            "wrong (high) — more reasoning effort did not reliably improve, and in this case actively "
            "hurt, accuracy. Small sample (n=4); treat as a pilot signal, not a settled finding."
        ),
    )
    callout.font = Font(italic=True)
    callout.alignment = WRAP
    ws.merge_cells(start_row=callout.row, start_column=1, end_row=callout.row, end_column=4)
    ws.row_dimensions[callout.row].height = 45

    for col, w in zip("ABCD", [14, 26, 20, 20]):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 12


# ── Ledger + main ────────────────────────────────────────────────────────────

def append_ledger(reports_dir: Path, entry: dict) -> None:
    ledger_path = reports_dir / "reports.jsonl"
    with ledger_path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(entry, sort_keys=True) + "\n")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--artifacts-root", default=".dist/study/artifacts_v2")
    ap.add_argument("--output", default=None, help="Override output path (default: auto under .dist/study/reports/<root>/model_comparison/)")
    ap.add_argument("--provider", required=True, help="Second model's provider (e.g. openai-compatible) — must match what wrote the tagged classification files.")
    ap.add_argument("--model", required=True, help="Second model's name (e.g. gpt-5-mini). Paired with --provider above.")
    ap.add_argument("--model-label", default=None, help="Display label for the second model in the report (default: --model value).")
    args = ap.parse_args()

    root = REPO_ROOT / args.artifacts_root
    root_name = Path(args.artifacts_root).name
    model_label = args.model_label or args.model

    records = load_records(root, args.provider, args.model)

    date_str = datetime.now().strftime("%Y-%m-%d")
    if args.output:
        out_file = REPO_ROOT / args.output
    else:
        suffix = model_slug(args.provider, args.model)
        out_file = REPO_ROOT / ".dist/study/reports" / root_name / "model_comparison" / f"ModelComparison_{date_str}.{suffix}.xlsx"
    out_file.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)
    build_comparison_tab(wb, records, model_label)
    build_accuracy_tab(wb, records, model_label)
    build_reasoning_effort_tab(wb, records, model_label)
    wb.save(out_file)

    reports_dir = REPO_ROOT / ".dist/study/reports" / root_name
    append_ledger(reports_dir, {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "git_sha": git_sha(),
        "report_type": "model_comparison",
        "bug_count": len(records),
        "provider": args.provider,
        "model": args.model,
        "output": str(out_file.relative_to(REPO_ROOT)),
    })

    print(f"wrote {out_file}: {len(records)} bugs")


if __name__ == "__main__":
    main()

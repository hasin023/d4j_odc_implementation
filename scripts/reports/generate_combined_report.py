"""Combined cross-project supervisor report: same 6-tab shape as the
per-project report (scripts/reports/generate_per_project_report.py), but
pooling all currently-collected projects into one workbook.

Rebuilt spec (2026-09-14):
- Tab 1 "Defect Types" and Tab 4 "Ablation Study" are flat bug-level tables
  with BOTH a Project and a Bug ID column (bug IDs collide across projects,
  e.g. every project has a "Bug 1"), sorted by (project, bug_id).
- Tab 2 "Type Distribution", Tab 3 "Type Drift", Tab 5/6 "Strategy Drift
  (Prefix/Postfix)" each contain one combined block ("All Projects") followed
  by one block per project below it — same matrix/table shape as the
  per-project report, just repeated per project. No paragraph-length
  description text anywhere; each block's only label is the project name
  (or "All Projects"), matching the old combined report's convention.

Output lands under .dist/study/reports/<artifacts_root_name>/combined/, and
each run appends one line to that root's reports.jsonl ledger.

Usage: python scripts/reports/generate_combined_report.py
       [--artifacts-root .dist/study/artifacts_v2] [--output <path>]
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT))
from d4j_odc_pipeline.odc import ODC_TYPE_NAMES, OTHER_TYPE_NAME  # noqa: E402

# Closure's full active-bug manifest is 174, but artifacts_v2 only has 152
# collected/classified so far (22 bugs pending or quarantined bad-fix-diff)
# — use the 152-bug subset manifest until the rest are collected.
MANIFESTS = {
    "Chart": ("manifest_chart26", 26),
    "Closure": ("manifest_closure152_v2", 152),
    "Lang": ("manifest_lang61", 61),
    "Math": ("manifest_math106", 106),
    "Mockito": ("manifest_mockito38", 38),
    "Time": ("manifest_time26", 26),
}
PROJECT_ORDER = ["Chart", "Closure", "Lang", "Math", "Mockito", "Time"]

HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_FONT = Font(bold=True)
HIGHLIGHT = PatternFill(start_color="FFF3B0", end_color="FFF3B0", fill_type="solid")
TITLE_FONT = Font(bold=True)

# Darker than HEADER_FILL so a block's title bar reads as a distinct section
# banner, not just another header row.
TITLE_FILL = PatternFill(start_color="8EA9DB", end_color="8EA9DB", fill_type="solid")
TITLE_BAR_FONT = Font(bold=True, color="000000")

# Blank rows between one block ("All Projects", "Chart", ...) and the next,
# used instead of a divider line to keep tables visually separated.
BLOCK_GAP_ROWS = 3

HEAT_LOW = (255, 255, 255)
HEAT_RED_HIGH = (230, 81, 0)     # deep orange/red — real drift/disagreement
HEAT_GREEN_HIGH = (46, 125, 50)  # deep green — diagonal, agreement


def git_sha() -> str | None:
    try:
        out = subprocess.run(["git", "rev-parse", "--short", "HEAD"], capture_output=True, text=True, timeout=5, cwd=REPO_ROOT)
    except (OSError, subprocess.SubprocessError):
        return None
    return out.stdout.strip() or None if out.returncode == 0 else None


def load_json(path: Path) -> dict | None:
    if not path.exists():
        return None
    return json.loads(path.read_text())


def load_classification(root: Path, project: str, bug_id: int, mode: str, tag: str) -> dict | None:
    return load_json(root / mode / f"{project}_{bug_id}_{mode}" / f"classification.{tag}.json")


def load_bug_url(root: Path, project: str, bug_id: int) -> str:
    ctx = load_json(root / "prefix" / f"{project}_{bug_id}_prefix" / "context.json")
    if not ctx:
        return "UNKNOWN"
    url = ctx.get("metadata", {}).get("report.url")
    if url and url != "UNKNOWN":
        return url
    bug_info = ctx.get("bug_info", "")
    m = re.search(r"Bug report url:\s*\n(.*)", bug_info)
    if m:
        return m.group(1).strip()
    return "UNKNOWN"


def set_url_cell(ws, row, col, url):
    cell = ws.cell(row=row, column=col, value=url)
    if url and url != "UNKNOWN":
        cell.hyperlink = url
        cell.font = Font(color="0563C1", underline="single")


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def style_title_bar(ws, row, ncols):
    """Fill a block's title row (e.g. "All Projects", "Lang") across the
    full width of the table below it, so it reads as one banner instead of
    a lone bold word in column A."""
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = TITLE_FILL
        cell.font = TITLE_BAR_FONT


def add_gap(ws, rows=BLOCK_GAP_ROWS):
    for _ in range(rows):
        ws.append([])


def heat_fill(value, max_value, high_rgb):
    if max_value <= 0 or value <= 0:
        return None
    t = 0.2 + 0.8 * (value / max_value)
    rgb = tuple(int(HEAT_LOW[i] + (high_rgb[i] - HEAT_LOW[i]) * t) for i in range(3))
    hexcolor = "%02X%02X%02X" % rgb
    return PatternFill(start_color=hexcolor, end_color=hexcolor, fill_type="solid")


def pct(n, total):
    return f"{(100.0 * n / total) if total else 0.0:.1f}%"


# ── Tab 1: Defect Types ─────────────────────────────────────────────────────

def build_defect_types_tab(wb, records):
    ws = wb.create_sheet("Defect Types")
    ws.append(["Project", "Bug ID", "URL", "Bug Type (Pre-fix)", "Bug Type (Post-fix)"])
    style_header(ws, 1, 5)
    for r in sorted(records, key=lambda r: (r["project"], r["bug_id"])):
        row = ws.max_row + 1
        ws.cell(row=row, column=1, value=r["project"])
        ws.cell(row=row, column=2, value=r["bug_id"])
        set_url_cell(ws, row, 3, r["url"])
        pre_v = r["sci_pre"] or "—"
        post_v = r["sci_post"] or "—"
        ws.cell(row=row, column=4, value=pre_v)
        ws.cell(row=row, column=5, value=post_v)
        if pre_v != post_v:
            for c in range(1, 6):
                ws.cell(row=row, column=c).fill = HIGHLIGHT
    widths = [10, 8, 45, 24, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── Tab 2: Type Distribution (All Projects, then per project) ───────────────

def write_distribution_block(ws, title, records, names):
    ws.append([title])
    style_title_bar(ws, ws.max_row, 5)
    ws.append(["ODC Type", "Pre-fix Count", "Pre-fix %", "Post-fix Count", "Post-fix %"])
    style_header(ws, ws.max_row, 5)
    total = len(records)
    counts_pre = Counter(r["sci_pre"] for r in records if r["sci_pre"] in names)
    counts_post = Counter(r["sci_post"] for r in records if r["sci_post"] in names)
    max_count = max(list(counts_pre.values()) + list(counts_post.values()) + [0])
    for name in names:
        pre_n = counts_pre.get(name, 0)
        post_n = counts_post.get(name, 0)
        ws.append([name, pre_n, pct(pre_n, total), post_n, pct(post_n, total)])
        row = ws.max_row
        fp = heat_fill(pre_n, max_count, HEAT_RED_HIGH)
        if fp:
            ws.cell(row=row, column=3).fill = fp
        fq = heat_fill(post_n, max_count, HEAT_RED_HIGH)
        if fq:
            ws.cell(row=row, column=5).fill = fq
    total_pre, total_post = sum(counts_pre.values()), sum(counts_post.values())
    ws.append(["Total", total_pre, pct(total_pre, total), total_post, pct(total_post, total)])
    for c in range(1, 6):
        ws.cell(row=ws.max_row, column=c).font = TITLE_FONT
    add_gap(ws)


def build_distribution_tab(wb, records, names):
    ws = wb.create_sheet("Type Distribution")
    write_distribution_block(ws, "All Projects", records, names)
    for project in PROJECT_ORDER:
        proj_records = [r for r in records if r["project"] == project]
        write_distribution_block(ws, project, proj_records, names)
    widths = [24, 14, 14, 14, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── Tab 3 / 5 / 6: confusion matrices (All Projects, then per project) ──────

def write_confusion_block(ws, title, records, key_row, key_col, names, row_label, col_label):
    ws.append([title])
    style_title_bar(ws, ws.max_row, len(names) + 1)
    ws.append([f"{row_label} \\ {col_label}"] + list(names))
    style_header(ws, ws.max_row, len(names) + 1)

    matrix = {p: Counter() for p in names}
    for r in records:
        p, q = r[key_row], r[key_col]
        if p in matrix and q in names:
            matrix[p][q] += 1

    off_diag_max = max((v for p in names for q in names if q != p for v in [matrix[p].get(q, 0)]), default=0)
    diag_max = max((matrix[p].get(p, 0) for p in names), default=0)

    for p in names:
        row_vals = [matrix[p].get(q, 0) for q in names]
        ws.append([p] + row_vals)
        row_idx = ws.max_row
        for j, q in enumerate(names, start=2):
            v = matrix[p].get(q, 0)
            if q == p:
                f = heat_fill(v, diag_max, HEAT_GREEN_HIGH)
            else:
                f = heat_fill(v, off_diag_max, HEAT_RED_HIGH)
            if f:
                ws.cell(row=row_idx, column=j).fill = f

    agreed = sum(matrix[p].get(p, 0) for p in names)
    total = len(records)
    ws.append(["Agreement Rate", f"{agreed}/{total}", pct(agreed, total)])
    style_header(ws, ws.max_row, 3)
    add_gap(ws)
    return agreed, total


def build_confusion_tab(wb, sheet_name, records, key_row, key_col, names, row_label, col_label):
    ws = wb.create_sheet(sheet_name)
    write_confusion_block(ws, "All Projects", records, key_row, key_col, names, row_label, col_label)
    for project in PROJECT_ORDER:
        proj_records = [r for r in records if r["project"] == project]
        write_confusion_block(ws, project, proj_records, key_row, key_col, names, row_label, col_label)
    ws.column_dimensions["A"].width = 26
    for i in range(2, len(names) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 16


def build_drift_tab(wb, records, names):
    build_confusion_tab(wb, "Type Drift", records, "sci_pre", "sci_post", names, "Pre-fix", "Post-fix")


def build_strategy_drift_tab(wb, sheet_name, records, key_sci, key_few, names):
    build_confusion_tab(wb, sheet_name, records, key_sci, key_few, names, "Scientific", "Few")


# ── Tab 4: Ablation Study ───────────────────────────────────────────────────

def build_ablation_tab(wb, records):
    ws = wb.create_sheet("Ablation Study")
    ws.append(["Project", "Bug ID", "Scientific: Prefix", "Few: Prefix", "Scientific: Postfix", "Few: Postfix"])
    style_header(ws, 1, 6)
    for r in sorted(records, key=lambda r: (r["project"], r["bug_id"])):
        row = ws.max_row + 1
        sci_pre, sci_post = r["sci_pre"] or "—", r["sci_post"] or "—"
        few_pre, few_post = r["few_pre"] or "—", r["few_post"] or "—"
        ws.cell(row=row, column=1, value=r["project"])
        ws.cell(row=row, column=2, value=r["bug_id"])
        ws.cell(row=row, column=3, value=sci_pre)
        ws.cell(row=row, column=4, value=few_pre)
        ws.cell(row=row, column=5, value=sci_post)
        ws.cell(row=row, column=6, value=few_post)
        if sci_pre != few_pre:
            ws.cell(row=row, column=3).fill = HIGHLIGHT
            ws.cell(row=row, column=4).fill = HIGHLIGHT
        if sci_post != few_post:
            ws.cell(row=row, column=5).fill = HIGHLIGHT
            ws.cell(row=row, column=6).fill = HIGHLIGHT
    widths = [10, 8, 20, 20, 20, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── Data loading ─────────────────────────────────────────────────────────────

def load_records(root: Path) -> list[dict]:
    records = []
    for project in PROJECT_ORDER:
        manifest_slug, expected = MANIFESTS[project]
        manifest = load_json(REPO_ROOT / f".dist/study/{manifest_slug}.json")
        bug_ids = sorted({e["bug_id"] for e in manifest["entries"]})
        assert len(bug_ids) == expected, f"{project}: expected {expected} bugs in manifest, got {len(bug_ids)}"
        for b in bug_ids:
            sci_pre = load_classification(root, project, b, "prefix", "scientific-open")
            sci_post = load_classification(root, project, b, "postfix", "scientific-open")
            few_pre = load_classification(root, project, b, "prefix", "few-open")
            few_post = load_classification(root, project, b, "postfix", "few-open")
            missing = [name for name, v in [
                ("scientific-open prefix", sci_pre), ("scientific-open postfix", sci_post),
                ("few-open prefix", few_pre), ("few-open postfix", few_post),
            ] if v is None]
            assert not missing, f"{project}-{b}: missing classification(s): {', '.join(missing)}"
            records.append({
                "project": project,
                "bug_id": b,
                "url": load_bug_url(root, project, b),
                "sci_pre": sci_pre.get("odc_type"),
                "sci_post": sci_post.get("odc_type"),
                "few_pre": few_pre.get("odc_type"),
                "few_post": few_post.get("odc_type"),
            })
    return records


def append_ledger(reports_dir: Path, entry: dict) -> None:
    ledger_path = reports_dir / "reports.jsonl"
    with ledger_path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(entry, sort_keys=True) + "\n")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--artifacts-root", default=".dist/study/artifacts_v2")
    ap.add_argument("--output", default=None, help="Override output path (default: auto under .dist/study/reports/<root>/combined/)")
    args = ap.parse_args()

    root = REPO_ROOT / args.artifacts_root
    root_name = Path(args.artifacts_root).name

    records = load_records(root)

    date_str = datetime.now().strftime("%Y-%m-%d")
    if args.output:
        out_file = REPO_ROOT / args.output
    else:
        out_file = REPO_ROOT / ".dist/study/reports" / root_name / "combined" / f"Combined_{date_str}.xlsx"
    out_file.parent.mkdir(parents=True, exist_ok=True)

    type_names = list(ODC_TYPE_NAMES) + [OTHER_TYPE_NAME]

    wb = Workbook()
    wb.remove(wb.active)
    build_defect_types_tab(wb, records)
    build_distribution_tab(wb, records, type_names)
    build_drift_tab(wb, records, type_names)
    build_ablation_tab(wb, records)
    build_strategy_drift_tab(wb, "Strategy Drift (Prefix)", records, "sci_pre", "few_pre", type_names)
    build_strategy_drift_tab(wb, "Strategy Drift (Postfix)", records, "sci_post", "few_post", type_names)
    wb.save(out_file)

    reports_dir = REPO_ROOT / ".dist/study/reports" / root_name
    append_ledger(reports_dir, {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "git_sha": git_sha(),
        "report_type": "combined",
        "projects": PROJECT_ORDER,
        "conditions": ["scientific-open", "few-open"],
        "bug_count": len(records),
        "output": str(out_file.relative_to(REPO_ROOT)),
    })

    print(f"wrote {out_file}: {len(records)} bugs across {len(PROJECT_ORDER)} projects")


if __name__ == "__main__":
    main()
